"""Streamlit UI for browsing and editing the condominium archive."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from importlib import import_module
import re
from typing import Any, Iterable, cast

from openpyxl import Workbook
from openpyxl.styles import Font
import streamlit as st
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, PropertySet, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from models import (
    Apartment,
    CondominiumArchive,
    Building,
    Condominium,
    Person,
    load_archive,
    save_archive,
)


st.set_page_config(
    page_title="Condominium Archive",
    page_icon="🏢",
    layout="wide",
    initial_sidebar_state="expanded",
)


def load_styles() -> None:
    """Inject the app's custom CSS (background, cards, badges, widget colors)."""
    st.markdown(
        """
        <style>
            .stApp {
                background:
                    radial-gradient(circle at top left, rgba(194, 220, 255, 0.95), transparent 25%),
                    radial-gradient(circle at 85% 0%, rgba(255, 227, 191, 0.85), transparent 22%),
                    linear-gradient(180deg, #f5efe3 0%, #edf2f7 100%);
            }
            header[data-testid="stHeader"] {
                background: transparent;
            }
            .block-container {
                padding-top: 0.8rem;
                padding-bottom: 2rem;
            }
            .hero-card {
                padding: 1.5rem 1.6rem;
                border-radius: 26px;
                background: rgba(255, 255, 255, 0.84);
                border: 1px solid rgba(19, 48, 70, 0.08);
                box-shadow: 0 20px 45px rgba(21, 37, 53, 0.08);
                backdrop-filter: blur(10px);
                margin-bottom: 1rem;
            }
            .eyebrow {
                font-size: 0.82rem;
                text-transform: uppercase;
                letter-spacing: 0.12em;
                color: #8a5a24;
                margin-bottom: 0.35rem;
                font-weight: 700;
            }
            .headline {
                font-size: 2.2rem;
                line-height: 1.05;
                font-weight: 800;
                color: #163047;
                margin: 0;
            }
            .subhead {
                color: #4a5f70;
                margin-top: 0.65rem;
                margin-bottom: 0;
                font-size: 1rem;
            }
            .card {
                padding: 1rem 1.1rem;
                border-radius: 20px;
                background: rgba(255, 255, 255, 0.78);
                border: 1px solid rgba(19, 48, 70, 0.08);
                box-shadow: 0 14px 30px rgba(21, 37, 53, 0.06);
            }
            .badge {
                display: inline-block;
                padding: 0.22rem 0.58rem;
                border-radius: 999px;
                background: #163047;
                color: white;
                font-size: 0.82rem;
                font-weight: 700;
                margin-bottom: 0.55rem;
            }
            button:disabled {
                opacity: 1 !important;
                cursor: default !important;
            }
            button[kind="secondary"]:disabled {
                background: linear-gradient(135deg, #163047 0%, #29506f 100%) !important;
                color: white !important;
                border: 1px solid rgba(22, 48, 71, 0.15) !important;
            }
            div[data-baseweb="select"] > div {
                background: linear-gradient(135deg, #163047 0%, #29506f 100%);
                border-color: rgba(22, 48, 71, 0.15);
                color: white;
            }
            div[data-baseweb="select"] span {
                color: white;
            }
            div[data-baseweb="select"] svg {
                fill: white;
            }
        </style>
        """,
        unsafe_allow_html=True,
    )


def get_archive() -> CondominiumArchive:
    """Load the archive into session state once, then reuse it across reruns."""
    if "archive" not in st.session_state:
        st.session_state.archive = load_archive()
    return st.session_state.archive


def get_active_section() -> str:
    pending_section = st.session_state.pop("next_active_section", None)
    if pending_section:
        st.session_state.active_section = pending_section
    if "active_section" not in st.session_state:
        st.session_state.active_section = "Condominiums"
    if st.session_state.active_section not in {
        "Condominiums",
        "Buildings",
        "Apartments",
    }:
        st.session_state.active_section = "Condominiums"
    return st.session_state.active_section


def save_and_refresh(
    archive: CondominiumArchive, message: str, active_section: str | None = None
) -> None:
    """Persist the archive to disk, queue a flash message, and rerun the app."""
    save_archive(archive)
    st.session_state.archive = archive
    st.session_state.flash_message = message
    if active_section:
        st.session_state.next_active_section = active_section
    st.rerun()


def show_flash_message() -> None:
    message = st.session_state.pop("flash_message", None)
    if message:
        st.success(message)


def format_currency(value: float) -> str:
    # Python formats with US separators (1,234.56); swap them to the European
    # convention (1.234,56) via a placeholder so the two substitutions don't collide.
    return f"€ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def total_registered_payments(transactions: Iterable) -> float:
    total = 0.0
    for transaction in transactions:
        if transaction.type == "payment":
            total += transaction.amount
    return round(total, 2)


def condominium_options(archive: CondominiumArchive) -> dict[str, Condominium]:
    ordered = sorted(archive.condominiums, key=lambda item: item.name.lower())
    return {condominium.name: condominium for condominium in ordered}


def building_options(condominium: Condominium) -> dict[str, Building]:
    ordered = sorted(condominium.buildings, key=lambda item: item.name.lower())
    return {building.name: building for building in ordered}


def apartment_options(building: Building) -> dict[str, Apartment]:
    ordered = sorted(building.apartments, key=lambda item: (item.unit, item.code))
    return {
        f"{apartment.code} · unit {apartment.unit} · floor {apartment.floor}": apartment
        for apartment in ordered
    }


def person_filled(first_name: str, last_name: str) -> bool:
    return bool(first_name.strip() and last_name.strip())


def build_person(
    *,
    first_name: str,
    last_name: str,
    phone: str = "",
    email: str = "",
    tax_code: str = "",
) -> Person:
    return Person(
        first_name=first_name.strip(),
        last_name=last_name.strip(),
        phone=phone.strip(),
        email=email.strip(),
        tax_code=tax_code.strip(),
    )


def slugify_filename(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "_", value.strip())
    return slug.strip("_") or "export"


def location_from_address(address: str) -> str:
    if not address.strip():
        return "Location not specified"
    parts = [part.strip() for part in address.split(",") if part.strip()]
    return parts[-1] if parts else address.strip()


def format_document_date(value: date) -> str:
    months = [
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    ]
    return f"{months[value.month - 1]} {value.day}, {value.year}"


def _pdf_paragraph(
    text: str, style: PropertySet, *, allow_markup: bool = False
) -> Paragraph:
    if allow_markup:
        safe_text = text.replace("&", "&amp;")
    else:
        safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(safe_text, cast(ParagraphStyle, style))


def _build_info_table(rows: list[tuple[str, str]], styles) -> Table:
    data = [
        [
            _pdf_paragraph(f"<b>{label}</b>", styles["BodyText"], allow_markup=True),
            _pdf_paragraph(value or "-", styles["BodyText"]),
        ]
        for label, value in rows
    ]
    table = Table(data, colWidths=[4.2 * cm, 11.8 * cm])
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _build_grid_table(headers: list[str], rows: list[list[str]], styles) -> Table:
    data = [
        [
            _pdf_paragraph(f"<b>{header}</b>", styles["BodyText"], allow_markup=True)
            for header in headers
        ]
    ]
    data.extend(
        [
            [_pdf_paragraph(value or "-", styles["BodyText"]) for value in row]
            for row in rows
        ]
    )
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9e7f5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#163047")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c9d4df")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f5f7fa")],
                ),
            ]
        )
    )
    return table


def build_component_pdf(
    *,
    title: str,
    info_rows: list[tuple[str, str]],
    movement_rows: list[list[str]],
    child_table: tuple[list[str], list[list[str]]] | None = None,
    footer_text: str | None = None,
) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.6 * cm,
        rightMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="SectionTitle",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#163047"),
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="FooterRight",
            parent=styles["BodyText"],
            alignment=2,
            textColor=colors.HexColor("#4a5f70"),
        )
    )

    story = [
        _pdf_paragraph(title, styles["Title"]),
        Spacer(1, 0.35 * cm),
        _pdf_paragraph("Main details", styles["SectionTitle"]),
        _build_info_table(info_rows, styles),
        Spacer(1, 0.35 * cm),
    ]

    if child_table and child_table[1]:
        headers, rows = child_table
        story.extend(
            [
                _pdf_paragraph("Linked items", styles["SectionTitle"]),
                _build_grid_table(headers, rows, styles),
                Spacer(1, 0.35 * cm),
            ]
        )

    story.append(_pdf_paragraph("Transactions", styles["SectionTitle"]))
    if movement_rows:
        story.append(
            _build_grid_table(
                ["Date", "Type", "Description", "Amount"], movement_rows, styles
            )
        )
    else:
        story.append(_pdf_paragraph("No transactions recorded.", styles["BodyText"]))

    if footer_text:
        story.extend(
            [
                Spacer(1, 0.5 * cm),
                _pdf_paragraph(footer_text, styles["FooterRight"]),
            ]
        )

    doc.build(story)
    return buffer.getvalue()


def _autosize_worksheet_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        values = [
            len(str(cell.value)) for cell in column_cells if cell.value is not None
        ]
        column_letter = column_cells[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(
            max(values, default=10) + 2,
            40,
        )


def _append_excel_table(worksheet, headers: list[str], rows: list[list[str]]) -> None:
    worksheet.append(headers)
    header_row = worksheet.max_row
    for cell in worksheet[header_row]:
        cell.font = Font(bold=True)

    if rows:
        for row in rows:
            worksheet.append(row)
    else:
        worksheet.append(["No data available."] + [""] * (len(headers) - 1))

    _autosize_worksheet_columns(worksheet)


def build_component_excel(
    *,
    title: str,
    info_rows: list[tuple[str, str]],
    movement_rows: list[list[str]],
    child_table: tuple[list[str], list[list[str]]] | None = None,
    footer_text: str | None = None,
) -> bytes:
    workbook = Workbook()
    info_sheet = workbook.active
    if info_sheet is None:
        info_sheet = workbook.create_sheet(title="Details")
    else:
        info_sheet.title = "Details"
    info_sheet.append([title])
    info_sheet["A1"].font = Font(bold=True)
    info_sheet.append([])
    info_sheet.append(["Main details"])
    info_sheet.cell(row=info_sheet.max_row, column=1).font = Font(bold=True)
    for label, value in info_rows:
        info_sheet.append([label, value or "-"])
        current_row = info_sheet.max_row
        info_sheet.cell(row=current_row, column=1).font = Font(bold=True)

    if child_table:
        child_headers, child_rows = child_table
        info_sheet.append([])
        info_sheet.append(["Linked items"])
        info_sheet.cell(row=info_sheet.max_row, column=1).font = Font(bold=True)
        _append_excel_table(info_sheet, child_headers, child_rows)

    info_sheet.append([])
    info_sheet.append(["Transactions"])
    info_sheet.cell(row=info_sheet.max_row, column=1).font = Font(bold=True)
    _append_excel_table(
        info_sheet,
        ["Date", "Type", "Description", "Amount"],
        movement_rows,
    )

    if footer_text:
        info_sheet.append([])
        info_sheet.append(["Location and date", footer_text])
        for cell in info_sheet[info_sheet.max_row]:
            if cell.column == 1:
                cell.font = Font(bold=True)
    _autosize_worksheet_columns(info_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_component_word(
    *,
    title: str,
    info_rows: list[tuple[str, str]],
    movement_rows: list[list[str]],
    child_table: tuple[list[str], list[list[str]]] | None = None,
    footer_text: str | None = None,
) -> bytes:
    try:
        docx_module = import_module("docx")
        wd_align_paragraph: Any = import_module("docx.enum.text").WD_ALIGN_PARAGRAPH
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "The 'python-docx' library is not installed. Install the package to export to Word."
        ) from error

    document = docx_module.Document()
    document.add_heading(title, level=1)

    document.add_heading("Main details", level=2)
    for label, value in info_rows:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = 0
        paragraph.paragraph_format.line_spacing = 1
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        paragraph.add_run(value or "-")

    if child_table and child_table[1]:
        child_headers, child_rows = child_table
        document.add_heading("Linked items", level=2)
        child_doc_table = document.add_table(rows=1, cols=len(child_headers))
        child_doc_table.style = "Table Grid"
        for index, header in enumerate(child_headers):
            child_doc_table.rows[0].cells[index].text = header
        for child_row in child_rows:
            row = child_doc_table.add_row().cells
            for index, value in enumerate(child_row):
                row[index].text = value or "-"

    document.add_heading("Transactions", level=2)
    if movement_rows:
        movement_table = document.add_table(rows=1, cols=4)
        movement_table.style = "Table Grid"
        movement_headers = ["Date", "Type", "Description", "Amount"]
        for index, header in enumerate(movement_headers):
            movement_table.rows[0].cells[index].text = header
        for movement_row in movement_rows:
            row = movement_table.add_row().cells
            for index, value in enumerate(movement_row):
                row[index].text = value or "-"
    else:
        document.add_paragraph("No transactions recorded.")

    if footer_text:
        document.add_paragraph("")
        footer_paragraph = document.add_paragraph(footer_text)
        footer_paragraph.alignment = wd_align_paragraph.RIGHT

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _close_export_dialog() -> None:
    st.session_state["close_export_dialog"] = True


def render_export_download_button(
    *, label: str, data: bytes, filename: str, mime: str
) -> None:
    st.download_button(
        label,
        data=data,
        file_name=filename,
        mime=mime,
        use_container_width=True,
        on_click=_close_export_dialog,
    )


@st.dialog("Export", width="large")
def open_export_dialog(*, title: str, exports: list[dict[str, str | bytes]]) -> None:
    if st.session_state.pop("close_export_dialog", False):
        st.rerun()
    st.caption(title)
    for export in exports:
        render_export_download_button(
            label=cast(str, export["label"]),
            data=cast(bytes, export["data"]),
            filename=cast(str, export["filename"]),
            mime=cast(str, export["mime"]),
        )


def persistent_selectbox(
    label: str,
    options: list[str],
    *,
    state_key: str,
    widget_key: str,
) -> str:
    """A selectbox whose selection survives reruns even if the widget is recreated."""
    if not options:
        raise ValueError("Selector options cannot be empty.")

    selected = st.session_state.get(state_key)
    if selected not in options:
        selected = options[0]
    st.session_state[state_key] = selected

    widget_value = st.session_state.get(widget_key)
    if widget_value not in options:
        st.session_state[widget_key] = selected

    def sync_selection() -> None:
        st.session_state[state_key] = st.session_state[widget_key]

    return st.selectbox(
        label,
        options,
        key=widget_key,
        on_change=sync_selection,
    )


def all_transactions(archive: CondominiumArchive) -> list:
    transactions = []
    for condominium in archive.condominiums:
        transactions.extend(condominium.transactions)
        for building in condominium.buildings:
            transactions.extend(building.transactions)
            for apartment in building.apartments:
                transactions.extend(apartment.transactions)
    return transactions


def render_header(archive: CondominiumArchive) -> None:
    st.markdown(
        f"""
        <div class="hero-card">
            <div class="eyebrow">Electronic Administrator</div>
            <h1 class="headline">AmEle</h1>
            <p class="subhead">
                {len(archive.condominiums)} condominiums, {archive.total_buildings()} buildings, and {archive.total_apartments()} apartments with separate statements.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_global_summary(archive: CondominiumArchive) -> None:
    all_payments = total_registered_payments(all_transactions(archive))
    occupants_assigned = sum(
        1
        for condominium in archive.condominiums
        for building in condominium.buildings
        for apartment in building.apartments
        if apartment.occupant
    )
    with st.sidebar:
        st.markdown("### Global overview")
        st.metric("Condominiums", len(archive.condominiums))
        st.metric("Buildings", archive.total_buildings())
        st.metric("Apartments", archive.total_apartments())
        st.metric("Total balance", format_currency(archive.total_balance()))
        st.metric("Assigned occupants", occupants_assigned)
        st.metric("Recorded payments", format_currency(all_payments))


def render_section_nav() -> str:
    sections = ["Condominiums", "Buildings", "Apartments"]
    active_section = get_active_section()
    cols = st.columns(len(sections))
    for col, section in zip(cols, sections):
        with col:
            clicked = st.button(
                section,
                key=f"nav_{section}",
                use_container_width=True,
                type="secondary",
                disabled=(section == active_section),
            )
            if clicked:
                st.session_state.active_section = section
                st.rerun()
    return active_section


def render_statement_panel(
    archive: CondominiumArchive,
    *,
    target,
    title: str,
    active_section: str,
    total_balance_label: str | None = None,
    total_balance_value: float | None = None,
) -> None:
    col1, col2 = st.columns([1.1, 1.2])
    with col1:
        st.write(f"**{title}**")
        st.metric("Own balance", format_currency(target.balance()))
        if total_balance_label and total_balance_value is not None:
            st.metric(
                total_balance_label,
                format_currency(total_balance_value),
            )
        if st.button(
            "New transaction",
            key=f"open_new_movement_{active_section}",
            use_container_width=True,
        ):
            open_add_transaction_dialog(
                archive,
                target=target,
                title=title,
                active_section=active_section,
            )

    with col2:
        st.write("**Recorded transactions**")
        if not target.transactions:
            st.caption("No transactions yet.")
        else:
            rows = [
                {
                    "Date": transaction.transaction_date,
                    "Type": transaction.type,
                    "Description": transaction.description,
                    "Amount": format_currency(transaction.amount),
                }
                for transaction in reversed(target.transactions)
            ]
            st.dataframe(rows, use_container_width=True, hide_index=True)


@st.dialog("New transaction", width="large")
def open_add_transaction_dialog(
    archive: CondominiumArchive,
    *,
    target,
    title: str,
    active_section: str,
) -> None:
    st.caption(title)
    with st.form(f"movement_form_{active_section}", clear_on_submit=True):
        description = st.text_input("Description")
        type_ = st.selectbox("Transaction type", ["charge", "payment"])
        amount = st.number_input("Amount", min_value=0.0, step=10.0, format="%.2f")
        submitted = st.form_submit_button(
            "Record transaction", use_container_width=True
        )
    if submitted:
        if not description.strip():
            st.error("Description is required.")
        else:
            target.add_transaction(
                description=description.strip(), amount=amount, type=type_
            )
            save_and_refresh(archive, "Transaction recorded.", active_section)


@st.dialog("Add condominium", width="large")
def open_add_condominium_dialog(archive: CondominiumArchive) -> None:
    with st.form("new_condominium", clear_on_submit=True):
        name = st.text_input("Condominium name")
        address = st.text_input("Address")
        notes = st.text_area("Notes", height=90)
        submitted = st.form_submit_button("Confirm add", use_container_width=True)
    if submitted:
        if not name.strip():
            st.error("Condominium name is required.")
        else:
            try:
                archive.add_condominium(
                    Condominium(
                        name=name.strip(),
                        address=address.strip(),
                        notes=notes.strip(),
                    )
                )
            except ValueError as error:
                st.error(str(error))
            else:
                save_and_refresh(archive, "Condominium added.", "Condominiums")


@st.dialog("Add building", width="large")
def open_add_building_dialog(
    archive: CondominiumArchive, condominium: Condominium
) -> None:
    st.caption(f"Selected condominium: {condominium.name}")
    with st.form("new_building", clear_on_submit=True):
        name = st.text_input("Building name")
        notes = st.text_area("Building notes", height=100)
        submitted = st.form_submit_button("Confirm add", use_container_width=True)
    if submitted:
        if not name.strip():
            st.error("Building name is required.")
        else:
            try:
                condominium.add_building(Building(name=name.strip(), notes=notes.strip()))
            except ValueError as error:
                st.error(str(error))
            else:
                save_and_refresh(archive, "Building added.", "Buildings")


@st.dialog("Add apartment", width="large")
def open_add_apartment_dialog(
    archive: CondominiumArchive,
    condominium: Condominium,
    building: Building,
) -> None:
    st.caption(
        f"Selected condominium: {condominium.name} | Selected building: {building.name}"
    )
    with st.form("new_apartment", clear_on_submit=True):
        c1, c2, c3 = st.columns(3)
        code = c1.text_input("Code")
        unit = c2.text_input("Unit")
        floor = c3.text_input("Floor")
        c4, c5, c6 = st.columns(3)
        area = c4.number_input("Area (sqm)", min_value=0.0, step=1.0, format="%.2f")
        shares = c5.number_input("Shares", min_value=0.0, step=1.0, format="%.2f")
        owner_first = c6.text_input("Owner first name")
        owner_cols = st.columns(3)
        owner_last = owner_cols[0].text_input("Owner last name")
        owner_phone = owner_cols[1].text_input("Owner phone")
        owner_email = owner_cols[2].text_input("Owner email")
        owner_tax_code = st.text_input("Owner tax code")
        notes = st.text_area("Notes", height=100)
        st.caption(
            "Optional occupant. If you leave first and last name empty, the owner will be used automatically."
        )
        o1, o2, o3 = st.columns(3)
        occ_first = o1.text_input("Occupant first name")
        occ_last = o2.text_input("Occupant last name")
        occ_phone = o3.text_input("Occupant phone")
        o4, o5 = st.columns(2)
        occ_email = o4.text_input("Occupant email")
        occ_tax_code = o5.text_input("Occupant tax code")
        submitted = st.form_submit_button("Confirm add", use_container_width=True)
    if submitted:
        if (
            not code.strip()
            or not unit.strip()
            or not floor.strip()
            or not person_filled(owner_first, owner_last)
        ):
            st.error(
                "Code, unit, floor, and the owner's first and last name are required."
            )
        else:
            owner = build_person(
                first_name=owner_first,
                last_name=owner_last,
                phone=owner_phone,
                email=owner_email,
                tax_code=owner_tax_code,
            )
            occupant = owner
            if occ_first.strip() or occ_last.strip():
                if not person_filled(occ_first, occ_last):
                    st.error(
                        "If you specify an occupant, first and last name are required."
                    )
                    return
                occupant = build_person(
                    first_name=occ_first,
                    last_name=occ_last,
                    phone=occ_phone,
                    email=occ_email,
                    tax_code=occ_tax_code,
                )
            try:
                building.add_apartment(
                    Apartment(
                        code=code.strip(),
                        unit=unit.strip(),
                        floor=floor.strip(),
                        area_sqm=round(area, 2),
                        shares=round(shares, 2),
                        owner=owner,
                        occupant=occupant,
                        notes=notes.strip(),
                    )
                )
            except ValueError as error:
                st.error(str(error))
            else:
                save_and_refresh(archive, "Apartment added.", "Apartments")


@st.dialog("Edit condominium", width="large")
def open_edit_condominium_dialog(
    archive: CondominiumArchive, target: Condominium
) -> None:
    with st.form("update_condominium"):
        name = st.text_input("Condominium name", value=target.name)
        address = st.text_input("Address", value=target.address)
        notes = st.text_area("Notes", value=target.notes, height=90)
        submitted = st.form_submit_button("Save changes", use_container_width=True)
    if submitted:
        if not name.strip():
            st.error("Condominium name is required.")
        else:
            try:
                archive.update_condominium(
                    target,
                    name=name,
                    address=address,
                    notes=notes,
                )
            except ValueError as error:
                st.error(str(error))
            else:
                save_and_refresh(archive, "Condominium updated.", "Condominiums")


@st.dialog("Edit building", width="large")
def open_edit_building_dialog(
    archive: CondominiumArchive, condominium: Condominium, target: Building
) -> None:
    st.caption(f"Selected condominium: {condominium.name}")
    with st.form("update_building"):
        name = st.text_input("Building name", value=target.name)
        notes = st.text_area("Building notes", value=target.notes, height=100)
        submitted = st.form_submit_button("Save changes", use_container_width=True)
    if submitted:
        if not name.strip():
            st.error("Building name is required.")
        else:
            try:
                condominium.update_building(target, name=name, notes=notes)
            except ValueError as error:
                st.error(str(error))
            else:
                save_and_refresh(archive, "Building updated.", "Buildings")


@st.dialog("Edit apartment", width="large")
def open_edit_apartment_dialog(
    archive: CondominiumArchive,
    building: Building,
    target: Apartment,
) -> None:
    with st.form("update_apartment"):
        c1, c2, c3 = st.columns(3)
        code = c1.text_input("Code", value=target.code)
        unit = c2.text_input("Unit", value=target.unit)
        floor = c3.text_input("Floor", value=target.floor)
        c4, c5, c6 = st.columns(3)
        area = c4.number_input(
            "Area (sqm)",
            min_value=0.0,
            step=1.0,
            format="%.2f",
            value=float(target.area_sqm),
        )
        shares = c5.number_input(
            "Shares",
            min_value=0.0,
            step=1.0,
            format="%.2f",
            value=float(target.shares),
        )
        owner_first = c6.text_input("Owner first name", value=target.owner.first_name)
        owner_cols = st.columns(3)
        owner_last = owner_cols[0].text_input(
            "Owner last name", value=target.owner.last_name
        )
        owner_phone = owner_cols[1].text_input(
            "Owner phone", value=target.owner.phone
        )
        owner_email = owner_cols[2].text_input(
            "Owner email", value=target.owner.email
        )
        owner_tax_code = st.text_input(
            "Owner tax code", value=target.owner.tax_code
        )
        notes = st.text_area("Notes", value=target.notes, height=100)
        st.caption(
            "Optional occupant. If you leave first and last name empty, the owner will be used automatically."
        )
        current_occupant = target.occupant
        occupant_matches_owner = (
            current_occupant == target.owner if current_occupant else False
        )
        occ_default_first = (
            "" if occupant_matches_owner else (current_occupant.first_name if current_occupant else "")
        )
        occ_default_last = (
            "" if occupant_matches_owner else (current_occupant.last_name if current_occupant else "")
        )
        occ_default_phone = (
            "" if occupant_matches_owner else (current_occupant.phone if current_occupant else "")
        )
        occ_default_email = (
            "" if occupant_matches_owner else (current_occupant.email if current_occupant else "")
        )
        occ_default_tax_code = (
            "" if occupant_matches_owner else (current_occupant.tax_code if current_occupant else "")
        )
        o1, o2, o3 = st.columns(3)
        occ_first = o1.text_input("Occupant first name", value=occ_default_first)
        occ_last = o2.text_input("Occupant last name", value=occ_default_last)
        occ_phone = o3.text_input("Occupant phone", value=occ_default_phone)
        o4, o5 = st.columns(2)
        occ_email = o4.text_input("Occupant email", value=occ_default_email)
        occ_tax_code = o5.text_input("Occupant tax code", value=occ_default_tax_code)
        submitted = st.form_submit_button("Save changes", use_container_width=True)
    if submitted:
        if (
            not code.strip()
            or not unit.strip()
            or not floor.strip()
            or not person_filled(owner_first, owner_last)
        ):
            st.error(
                "Code, unit, floor, and the owner's first and last name are required."
            )
        else:
            owner = build_person(
                first_name=owner_first,
                last_name=owner_last,
                phone=owner_phone,
                email=owner_email,
                tax_code=owner_tax_code,
            )
            occupant = owner
            if occ_first.strip() or occ_last.strip():
                if not person_filled(occ_first, occ_last):
                    st.error(
                        "If you specify an occupant, first and last name are required."
                    )
                    return
                occupant = build_person(
                    first_name=occ_first,
                    last_name=occ_last,
                    phone=occ_phone,
                    email=occ_email,
                    tax_code=occ_tax_code,
                )
            try:
                building.update_apartment(
                    target,
                    code=code,
                    unit=unit,
                    floor=floor,
                    area_sqm=area,
                    shares=shares,
                    owner=owner,
                    notes=notes,
                )
                target.occupant = occupant
            except ValueError as error:
                st.error(str(error))
            else:
                save_and_refresh(archive, "Apartment updated.", "Apartments")


@st.dialog("Confirm deletion", width="large")
def open_delete_condominium_dialog(
    archive: CondominiumArchive, condominium_name: str
) -> None:
    st.warning(
        f"You are about to delete the condominium '{condominium_name}' along with all its buildings, apartments, and statements."
    )
    col1, col2 = st.columns(2)
    with col1:
        if st.button(
            "Confirm deletion",
            key="confirm_delete_condominium",
            use_container_width=True,
        ):
            archive.remove_condominium(condominium_name)
            save_and_refresh(archive, "Condominium removed.", "Condominiums")
    with col2:
        if st.button(
            "Cancel",
            key="cancel_delete_condominium",
            use_container_width=True,
            type="secondary",
        ):
            st.rerun()


@st.dialog("Confirm deletion", width="large")
def open_delete_building_dialog(
    archive: CondominiumArchive,
    condominium: Condominium,
    building_name: str,
) -> None:
    st.warning(
        f"You are about to delete the building '{building_name}' of condominium '{condominium.name}' along with all linked apartments and statements."
    )
    col1, col2 = st.columns(2)
    with col1:
        if st.button(
            "Confirm deletion",
            key="confirm_delete_building",
            use_container_width=True,
        ):
            condominium.remove_building(building_name)
            save_and_refresh(archive, "Building removed.", "Buildings")
    with col2:
        if st.button(
            "Cancel",
            key="cancel_delete_building",
            use_container_width=True,
            type="secondary",
        ):
            st.rerun()


@st.dialog("Confirm deletion", width="large")
def open_delete_apartment_dialog(
    archive: CondominiumArchive,
    building: Building,
    apartment_code: str,
) -> None:
    st.warning(
        f"You are about to delete apartment '{apartment_code}' and its statement."
    )
    col1, col2 = st.columns(2)
    with col1:
        if st.button(
            "Confirm deletion",
            key="confirm_delete_apartment",
            use_container_width=True,
        ):
            building.remove_apartment(apartment_code)
            save_and_refresh(archive, "Apartment removed.", "Apartments")
    with col2:
        if st.button(
            "Cancel",
            key="cancel_delete_apartment",
            use_container_width=True,
            type="secondary",
        ):
            st.rerun()


def render_condominium_tab(archive: CondominiumArchive) -> None:
    if not archive.condominiums:
        st.info("Add the first condominium to get started.")

    if st.button("Add condominium", key="open_add_condominium"):
        open_add_condominium_dialog(archive)

    target = None
    if archive.condominiums:
        options = condominium_options(archive)
        target_name = persistent_selectbox(
            "Select condominium",
            list(options.keys()),
            state_key="state_manage_condominium",
            widget_key="manage_condominium",
        )
        target = options[target_name]
        col1, col2, col3 = st.columns(3)
        with col1:
            if st.button(
                "Edit condominium",
                use_container_width=True,
                key="edit_condominium_button",
            ):
                open_edit_condominium_dialog(archive, target)
        with col2:
            if st.button(
                "Delete condominium",
                type="secondary",
                use_container_width=True,
                key="delete_condominium_button",
            ):
                open_delete_condominium_dialog(archive, target_name)
        with col3:
            if st.button(
                "Export condominium",
                use_container_width=True,
                key="export_condominium_button",
            ):
                export_title = f"Condominium {target.name}"
                footer_text = (
                    f"{location_from_address(target.address)}, "
                    f"{format_document_date(date.today())}"
                )
                info_rows = [
                    ("Name", target.name),
                    ("Address", target.address or "-"),
                    ("Notes", target.notes or "-"),
                    ("Own balance", format_currency(target.balance())),
                    ("Total balance", format_currency(target.total_balance())),
                ]
                child_table = (
                    [
                        "Building",
                        "Apartments",
                        "Own balance",
                        "Total balance",
                    ],
                    [
                        [
                            building.name,
                            str(len(building.apartments)),
                            format_currency(building.balance()),
                            format_currency(building.total_balance()),
                        ]
                        for building in target.buildings
                    ],
                )
                movement_rows = [
                    [
                        transaction.transaction_date,
                        transaction.type,
                        transaction.description,
                        format_currency(transaction.amount),
                    ]
                    for transaction in target.transactions
                ]
                slug_name = slugify_filename(target.name)
                open_export_dialog(
                    title=export_title,
                    exports=[
                        {
                            "label": "Download PDF",
                            "data": build_component_pdf(
                                title=export_title,
                                info_rows=info_rows,
                                child_table=child_table,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": f"condominium_{slug_name}.pdf",
                            "mime": "application/pdf",
                        },
                        {
                            "label": "Download Word",
                            "data": build_component_word(
                                title=export_title,
                                info_rows=info_rows,
                                child_table=child_table,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": f"condominium_{slug_name}.docx",
                            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        },
                        {
                            "label": "Download Excel",
                            "data": build_component_excel(
                                title=export_title,
                                info_rows=info_rows,
                                child_table=child_table,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": f"condominium_{slug_name}.xlsx",
                            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        },
                    ],
                )
        st.caption(
            "Deleting also removes the linked buildings, apartments, and statements."
        )

    if not archive.condominiums or target is None:
        return

    rows = []
    for condominium in archive.condominiums:
        rows.append(
            {
                "Condominium": condominium.name,
                "Address": condominium.address or "-",
                "Buildings": len(condominium.buildings),
                "Apartments": condominium.total_apartments(),
                "Own balance": format_currency(condominium.balance()),
                "Total balance": format_currency(condominium.total_balance()),
            }
        )
    st.dataframe(rows, use_container_width=True, hide_index=True)
    render_statement_panel(
        archive,
        target=target,
        title=f"Statement for condominium {target.name}",
        active_section="Condominiums",
        total_balance_label="Total balance including child levels",
        total_balance_value=target.total_balance(),
    )


def render_buildings_tab(archive: CondominiumArchive) -> None:
    if not archive.condominiums:
        st.info("You must first create at least one condominium.")
        return

    condominium_map = condominium_options(archive)
    selected_name = persistent_selectbox(
        "Condominium",
        list(condominium_map.keys()),
        state_key="state_building_condominium",
        widget_key="building_condominium",
    )
    condominium = condominium_map[selected_name]

    if st.button("Add building", key="open_add_building"):
        open_add_building_dialog(archive, condominium)

    if not condominium.buildings:
        st.caption("No buildings available.")
    else:
        options = building_options(condominium)
        target_name = persistent_selectbox(
            "Select building",
            list(options.keys()),
            state_key="state_manage_building",
            widget_key="manage_building",
        )
        target = options[target_name]
        col1, col2, col3 = st.columns(3)
        with col1:
            if st.button(
                "Edit building",
                use_container_width=True,
                key="edit_building_button",
            ):
                open_edit_building_dialog(archive, condominium, target)
        with col2:
            if st.button(
                "Delete building",
                type="secondary",
                use_container_width=True,
                key="delete_building_button",
            ):
                open_delete_building_dialog(archive, condominium, target_name)
        with col3:
            if st.button(
                "Export building",
                use_container_width=True,
                key="export_building_button",
            ):
                export_title = f"Building {target.name}"
                footer_text = (
                    f"{location_from_address(condominium.address)}, "
                    f"{format_document_date(date.today())}"
                )
                info_rows = [
                    ("Condominium", condominium.name),
                    ("Name", target.name),
                    ("Notes", target.notes or "-"),
                    ("Own balance", format_currency(target.balance())),
                    ("Total balance", format_currency(target.total_balance())),
                ]
                child_table = (
                    ["Code", "Unit", "Floor", "Owner", "Own balance"],
                    [
                        [
                            apartment.code,
                            apartment.unit,
                            apartment.floor,
                            apartment.owner.full_name,
                            format_currency(apartment.balance()),
                        ]
                        for apartment in target.apartments
                    ],
                )
                movement_rows = [
                    [
                        transaction.transaction_date,
                        transaction.type,
                        transaction.description,
                        format_currency(transaction.amount),
                    ]
                    for transaction in target.transactions
                ]
                slug_condominium = slugify_filename(condominium.name)
                slug_building = slugify_filename(target.name)
                open_export_dialog(
                    title=export_title,
                    exports=[
                        {
                            "label": "Download PDF",
                            "data": build_component_pdf(
                                title=export_title,
                                info_rows=info_rows,
                                child_table=child_table,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": (
                                f"building_{slug_condominium}_{slug_building}.pdf"
                            ),
                            "mime": "application/pdf",
                        },
                        {
                            "label": "Download Word",
                            "data": build_component_word(
                                title=export_title,
                                info_rows=info_rows,
                                child_table=child_table,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": (
                                f"building_{slug_condominium}_{slug_building}.docx"
                            ),
                            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        },
                        {
                            "label": "Download Excel",
                            "data": build_component_excel(
                                title=export_title,
                                info_rows=info_rows,
                                child_table=child_table,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": (
                                f"building_{slug_condominium}_{slug_building}.xlsx"
                            ),
                            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        },
                    ],
                )
        st.caption(
            "Deleting also removes the building's apartments and statements."
        )
        render_statement_panel(
            archive,
            target=target,
            title=f"Statement for building {target.name}",
            active_section="Buildings",
            total_balance_label="Total balance including apartments",
            total_balance_value=target.total_balance(),
        )

    rows = [
        {
            "Building": building.name,
            "Apartments": len(building.apartments),
            "Own balance": format_currency(building.balance()),
            "Total balance": format_currency(building.total_balance()),
        }
        for building in condominium.buildings
    ]
    if rows:
        st.dataframe(rows, use_container_width=True, hide_index=True)


def render_apartments_tab(archive: CondominiumArchive) -> None:
    if not archive.condominiums:
        st.info("You must first create at least one condominium.")
        return

    condominium_map = condominium_options(archive)
    selected_condominium = persistent_selectbox(
        "Condominium",
        list(condominium_map.keys()),
        state_key="state_apartment_condominium",
        widget_key="apartment_condominium",
    )
    condominium = condominium_map[selected_condominium]
    if not condominium.buildings:
        st.info("The selected condominium has no buildings yet.")
        return

    building_map = building_options(condominium)
    selected_building = persistent_selectbox(
        "Building",
        list(building_map.keys()),
        state_key="state_apartment_building",
        widget_key="apartment_building",
    )
    building = building_map[selected_building]

    if st.button("Add apartment", key="open_add_apartment"):
        open_add_apartment_dialog(archive, condominium, building)

    target: Apartment | None = None
    if not building.apartments:
        st.caption("No apartments available.")
    else:
        options = apartment_options(building)
        target_label = persistent_selectbox(
            "Select apartment",
            list(options.keys()),
            state_key="state_manage_apartment",
            widget_key="manage_apartment",
        )
        target = options[target_label]
        col1, col2, col3 = st.columns(3)
        with col1:
            if st.button(
                "Edit apartment",
                use_container_width=True,
                key="edit_apartment_button",
            ):
                open_edit_apartment_dialog(archive, building, target)
        with col2:
            if st.button(
                "Delete apartment",
                type="secondary",
                use_container_width=True,
                key="delete_apartment_button",
            ):
                open_delete_apartment_dialog(archive, building, target.code)
        with col3:
            if st.button(
                "Export apartment",
                use_container_width=True,
                key="export_apartment_button",
            ):
                export_title = f"Apartment {target.code}"
                footer_text = (
                    f"{location_from_address(condominium.address)}, "
                    f"{format_document_date(date.today())}"
                )
                info_rows = [
                    ("Condominium", condominium.name),
                    ("Building", building.name),
                    ("Code", target.code),
                    ("Unit", target.unit),
                    ("Floor", target.floor),
                    ("Area", f"{target.area_sqm:.2f} sqm"),
                    ("Shares", f"{target.shares:.2f}"),
                    ("Owner", target.owner.full_name),
                    ("Owner phone", target.owner.phone or "-"),
                    ("Owner email", target.owner.email or "-"),
                    (
                        "Owner tax code",
                        target.owner.tax_code or "-",
                    ),
                    (
                        "Occupant",
                        target.occupant.full_name if target.occupant else "-",
                    ),
                    (
                        "Occupant phone",
                        (
                            target.occupant.phone
                            if target.occupant and target.occupant.phone
                            else "-"
                        ),
                    ),
                    (
                        "Occupant email",
                        (
                            target.occupant.email
                            if target.occupant and target.occupant.email
                            else "-"
                        ),
                    ),
                    (
                        "Occupant tax code",
                        (
                            target.occupant.tax_code
                            if target.occupant and target.occupant.tax_code
                            else "-"
                        ),
                    ),
                    ("Notes", target.notes or "-"),
                    ("Own balance", format_currency(target.balance())),
                ]
                movement_rows = [
                    [
                        transaction.transaction_date,
                        transaction.type,
                        transaction.description,
                        format_currency(transaction.amount),
                    ]
                    for transaction in target.transactions
                ]
                slug_condominium = slugify_filename(condominium.name)
                slug_building = slugify_filename(building.name)
                slug_apartment = slugify_filename(target.code)
                open_export_dialog(
                    title=export_title,
                    exports=[
                        {
                            "label": "Download PDF",
                            "data": build_component_pdf(
                                title=export_title,
                                info_rows=info_rows,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": (
                                f"apartment_{slug_condominium}_{slug_building}_{slug_apartment}.pdf"
                            ),
                            "mime": "application/pdf",
                        },
                        {
                            "label": "Download Word",
                            "data": build_component_word(
                                title=export_title,
                                info_rows=info_rows,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": (
                                f"apartment_{slug_condominium}_{slug_building}_{slug_apartment}.docx"
                            ),
                            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        },
                        {
                            "label": "Download Excel",
                            "data": build_component_excel(
                                title=export_title,
                                info_rows=info_rows,
                                movement_rows=movement_rows,
                                footer_text=footer_text,
                            ),
                            "filename": (
                                f"apartment_{slug_condominium}_{slug_building}_{slug_apartment}.xlsx"
                            ),
                            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        },
                    ],
                )
        st.caption("Deleting also removes the apartment's statement.")

    rows = [
        {
            "Code": apartment.code,
            "Unit": apartment.unit,
            "Floor": apartment.floor,
            "Owner": apartment.owner.full_name,
            "Occupant": (
                apartment.occupant.full_name if apartment.occupant else "Not assigned"
            ),
            "Own balance": format_currency(apartment.balance()),
        }
        for apartment in building.apartments
    ]
    if rows and target is not None:
        st.dataframe(rows, use_container_width=True, hide_index=True)
        render_statement_panel(
            archive,
            target=target,
            title=f"Statement for apartment {target.code}",
            active_section="Apartments",
        )


def main() -> None:
    load_styles()
    archive = get_archive()
    get_active_section()
    render_global_summary(archive)
    show_flash_message()
    render_header(archive)
    active_section = render_section_nav()

    if active_section == "Condominiums":
        render_condominium_tab(archive)
    elif active_section == "Buildings":
        render_buildings_tab(archive)
    else:
        render_apartments_tab(archive)


if __name__ == "__main__":
    main()
